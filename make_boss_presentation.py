#!/usr/bin/env python3
"""Build a concise four-slide English executive update."""

import json
from pathlib import Path

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from flow_cytometry.config import Config
from flow_cytometry.data import SYNTHETIC_CLASS_NAMES, ToyCells


ROOT = Path(__file__).resolve().parent
TRAINED_DIR = ROOT / "flow_cytometry_trained"
ASSET_DIR = ROOT / "presentation_assets"
OUTPUT_PATH = ROOT / "coherent_flow_cytometry_update.pptx"

BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
LIGHT_GRAY = RGBColor(244, 244, 244)
LIGHT_BLUE = RGBColor(224, 236, 247)
MID_GRAY = RGBColor(100, 100, 100)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    size: float = 20,
    bold: bool = False,
    color: RGBColor = BLACK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
) -> None:
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    add_text(slide, title, 0.60, 0.25, 12.1, 0.55, size=28, bold=True)
    divider = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.60),
        Inches(0.86),
        Inches(12.1),
        Inches(0.025),
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = BLACK
    divider.line.fill.background()
    if subtitle:
        add_text(slide, subtitle, 0.62, 0.96, 12.0, 0.36, size=15, color=MID_GRAY)


def add_box(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    title: str,
    body: str = "",
    *,
    fill_color: RGBColor = LIGHT_GRAY,
    title_size: float = 19,
    body_size: float = 16,
    centered: bool = False,
) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = BLACK
    shape.line.width = Pt(1)
    alignment = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT
    add_text(
        slide,
        title,
        x + 0.15,
        y + 0.10,
        width - 0.30,
        0.42,
        size=title_size,
        bold=True,
        align=alignment,
    )
    if body:
        add_text(
            slide,
            body,
            x + 0.15,
            y + 0.56,
            width - 0.30,
            height - 0.62,
            size=body_size,
            align=alignment,
        )


def add_footer(slide, number: int, text: str = "") -> None:
    if text:
        add_text(slide, text, 0.60, 7.10, 11.6, 0.18, size=7.5, color=MID_GRAY)
    add_text(slide, str(number), 12.25, 7.04, 0.40, 0.22, size=9, color=MID_GRAY, align=PP_ALIGN.RIGHT)


def set_speaker_notes(slide, text: str) -> None:
    """Write concise presenter guidance into the slide's notes pane."""
    for placeholder in slide.notes_slide.placeholders:
        if placeholder.placeholder_format.type == 2:  # BODY
            placeholder.text_frame.text = text
            return
    raise RuntimeError("Speaker-notes body placeholder was not found")


def create_assets() -> tuple[Path, Path]:
    """Create simplified cell and accuracy graphics for the compact deck."""
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    config = Config()
    dataset = ToyCells(2, config, config.seed)
    cells_path = ASSET_DIR / "cell_phase_classes.png"
    fig, axes = plt.subplots(1, 4, figsize=(12.8, 3.0), constrained_layout=True)
    image = None
    for label, axis in enumerate(axes):
        phase = next(value for value, item_label in dataset if item_label == label)
        image = axis.imshow(
            phase,
            cmap="viridis",
            vmin=0,
            vmax=config.phase_depth_rad,
            extent=(-1, 1, -1, 1),
        )
        axis.set_title(
            f"Class {label}: {SYNTHETIC_CLASS_NAMES[label]}", fontsize=15
        )
        axis.set_xticks([])
        axis.set_yticks([])
    fig.colorbar(image, ax=axes, label="Phase [rad]", fraction=0.025)
    fig.savefig(cells_path, dpi=220, facecolor="white", bbox_inches="tight")
    plt.close(fig)

    metrics = json.loads((TRAINED_DIR / "metrics.json").read_text(encoding="utf-8"))
    accuracy = [100 * value for value in metrics["history"]["accuracy"]]
    accuracy_path = ASSET_DIR / "accuracy_curve_compact.png"
    fig, axis = plt.subplots(figsize=(6.4, 2.65))
    axis.plot(range(1, 9), accuracy, "o-", color="black", lw=2.7)
    axis.set_xlim(1, 8)
    axis.set_ylim(78, 101.5)
    axis.set_xlabel("Epoch", fontsize=13)
    axis.set_ylabel("Test accuracy [%]", fontsize=13)
    axis.set_title("100% from epoch 2", fontsize=16, fontweight="bold")
    axis.grid(axis="y", alpha=0.2)
    axis.spines[["top", "right"]].set_visible(False)
    axis.tick_params(labelsize=11)
    fig.tight_layout()
    fig.savefig(accuracy_path, dpi=220, facecolor="white", bbox_inches="tight")
    plt.close(fig)
    return cells_path, accuracy_path


def build_deck(cells_path: Path, accuracy_path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    # 1. Motivation and executive takeaway.
    slide = presentation.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    add_title(
        slide,
        "Coherent Meta-Optical Flow Cytometry",
        "Paper-inspired end-to-end optimization of optics and digital classification",
    )
    add_box(
        slide,
        0.75,
        1.55,
        3.80,
        2.20,
        "Paper insight",
        "Coherent propagation mixes the optical field so class information survives a small detector.",
        fill_color=LIGHT_BLUE,
        title_size=21,
        body_size=18,
    )
    add_box(
        slide,
        4.77,
        1.55,
        3.80,
        2.20,
        "Our adaptation",
        "A flowing phase cell sits between two trainable phase masks at each station.",
        title_size=21,
        body_size=18,
    )
    add_box(
        slide,
        8.79,
        1.55,
        3.80,
        2.20,
        "Why it matters",
        "Optical preprocessing compresses the signal before ADC and the digital backend.",
        fill_color=LIGHT_BLUE,
        title_size=21,
        body_size=18,
    )
    add_box(slide, 0.95, 4.45, 3.35, 1.25, "3 stations", "M1 → cell → M2", centered=True, title_size=23, body_size=18)
    add_box(slide, 4.98, 4.45, 3.35, 1.25, "48 features", "3 × 4 × 4 detector values", centered=True, title_size=23, body_size=18)
    add_box(slide, 9.01, 4.45, 3.35, 1.25, "100%", "64 / 64 synthetic test cells", centered=True, fill_color=LIGHT_BLUE, title_size=27, body_size=17)
    add_text(slide, "Goal: move useful feature extraction into passive optics while keeping the electronic classifier small.", 1.00, 6.25, 11.3, 0.48, size=18, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 1, "Method basis: N. Behrens, Y. Li, F. Monticone, arXiv:2608.10304 (2026).")
    set_speaker_notes(
        slide,
        "Motivation: perform classification-relevant preprocessing before digitization. "
        "Behrens, Li, and Monticone show that coherent propagation can reshape joint "
        "sensor statistics under a strong detector bottleneck. We adapt this principle "
        "to flow cytometry using two trainable phase masks around the cell at each "
        "station. The optics and the classifier are optimized together. The reported "
        "100% is a synthetic proof-of-concept result, not biological accuracy.",
    )

    # 2. Optical setup and parameters.
    slide = presentation.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    add_title(slide, "Optical setup", "Paired phase-only metasurfaces repeat along a 1 m/s microfluidic channel.")
    slide.shapes.add_picture(
        str(TRAINED_DIR / "optical_setup.png"),
        Inches(0.45),
        Inches(1.25),
        width=Inches(8.55),
    )
    add_text(slide, "KEY PARAMETERS", 9.20, 1.30, 3.5, 0.40, size=20, bold=True)
    add_text(
        slide,
        "λ = 632.8 nm\n\nM1 → cell: 8 mm\ncell → M2: 8 mm\nM2 → sensor: 4 mm\n\nGrid: 192 × 192\ndx = 0.35 µm\nFOV = 67.2 µm\n\nMasks: 32 × 32 controls\nPhase range: −π to +π\n\n3 stations simulated\n10–100 target",
        9.20,
        1.75,
        3.55,
        4.95,
        size=17,
    )
    add_footer(slide, 2, "Paper reference: 1200 × 1200 grid, 0.35 µm pitch, 420 µm aperture, and 16 mm + 4 mm propagation.")
    set_speaker_notes(
        slide,
        "Method: a 632.8 nm coherent plane wave illuminates each station. M1 creates "
        "structured illumination; the field propagates 8 mm to the thin phase cell, "
        "8 mm to M2, and 4 mm to the sensor. Propagation is simulated with the angular "
        "spectrum method. We currently simulate three stations, with two 32 by 32 "
        "phase-control grids per station, interpolated to the 192 by 192 simulation grid.",
    )

    # 3. Learned masks and detector priorities.
    slide = presentation.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    add_title(slide, "Learned phase masks and detector requirements")
    slide.shapes.add_picture(
        str(TRAINED_DIR / "phase_masks.png"),
        Inches(0.65),
        Inches(1.15),
        height=Inches(5.80),
    )
    add_box(slide, 5.15, 1.30, 2.25, 1.25, "16.8 µm", "effective pixel pitch", centered=True, fill_color=LIGHT_BLUE, title_size=24, body_size=14)
    add_box(slide, 7.68, 1.30, 2.35, 1.25, "67.2 × 67.2", "µm² active area", centered=True, title_size=22, body_size=14)
    add_box(slide, 10.31, 1.30, 2.25, 1.25, "≥59.5 kfps", "continuous target", centered=True, fill_color=LIGHT_BLUE, title_size=22, body_size=14)
    add_box(slide, 5.15, 2.95, 7.41, 1.35, "Detector readout", "4 × 4 pixels / station × 3 stations = 48 intensity features", title_size=21, body_size=18)
    add_box(slide, 5.15, 4.60, 7.41, 1.55, "Timing at 1 m/s", "67.2 µs transit • ≥14.9 kfps for one frame/FOV\nExposure ≤16.8 µs limits motion to one effective pixel", fill_color=LIGHT_BLUE, title_size=21, body_size=17)
    add_text(slide, "M1: structured illumination     M2: coherent mixing / compression     phase range: ±π", 5.25, 6.43, 7.15, 0.38, size=15, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 3, "Current detector model is ideal: 100% fill factor and no noise, quantization, or saturation.")
    set_speaker_notes(
        slide,
        "The 4 by 4 detector is an intentional information bottleneck; the paper predicts "
        "the largest optical benefit in this regime. Each detector pixel averages 48 by "
        "48 simulation samples, giving a 16.8 micrometer effective pitch and a 67.2 "
        "micrometer square active area. At 1 meter per second, transit time is 67.2 "
        "microseconds. A 59.5 kiloframe-per-second target with exposure below 16.8 "
        "microseconds limits motion to one effective pixel. Detector noise is not yet modeled.",
    )

    # 4. Cells and performance.
    slide = presentation.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    add_title(slide, "Synthetic phase cells and classification performance")
    slide.shapes.add_picture(
        str(cells_path), Inches(0.55), Inches(1.08), width=Inches(12.15)
    )
    slide.shapes.add_picture(
        str(accuracy_path), Inches(0.75), Inches(4.05), width=Inches(5.55)
    )
    add_box(slide, 6.70, 4.18, 2.55, 1.45, "100%", "64 / 64 test cells", centered=True, fill_color=LIGHT_BLUE, title_size=29, body_size=17)
    add_box(slide, 9.55, 4.18, 2.75, 1.45, "8 epochs", "joint M1 + M2 + linear training", centered=True, title_size=25, body_size=15)
    add_text(slide, "Next: measured cell phase maps + detector noise + fabrication tolerances", 6.65, 6.05, 5.75, 0.52, size=18, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 4, "Proof of concept only: deliberately separable synthetic morphology, not biological accuracy.")
    set_speaker_notes(
        slide,
        "Training uses four synthetic phase morphologies with independent position, "
        "rotation, and size variation: 192 training cells and 64 held-out test cells. "
        "The three stations produce 48 normalized intensity features, and a single "
        "linear layer maps them to four classes. AdamW and cross-entropy jointly train "
        "M1, M2, and the classifier for eight epochs. Accuracy reaches 100% from epoch "
        "two. The next milestone is realism: measured cell phases, detector noise, and "
        "fabrication-tolerance testing.",
    )

    presentation.save(OUTPUT_PATH)


def main() -> None:
    cells_path, accuracy_path = create_assets()
    build_deck(cells_path, accuracy_path)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
